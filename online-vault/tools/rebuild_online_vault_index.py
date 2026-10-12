#!/usr/bin/env python3
"""Rebuild the unified online-vault index.

The script keeps source vaults untouched and writes generated indexes/pages:
- online-vault-index.json / online-vault-index.js
- index.html
- gemini-vault/search-index.json
- gemini-vault/html/search-index.js
- gemini-vault/html/index.html
- gemini-vault/html/items/*.html
- logs/online-vault-summary.json
- gemini-vault/logs/reindex-summary.json
"""

from __future__ import annotations

import csv
import hashlib
import html as html_lib
import json
import re
import zipfile
from collections import Counter, defaultdict
from datetime import datetime
from html.parser import HTMLParser
from pathlib import Path
from typing import Any
from urllib.parse import quote


ROOT = Path(__file__).resolve().parents[1]
WECHAT_DIR = ROOT / "wechat-vault"
WEB_DIR = ROOT / "web-vault"
GEMINI_DIR = ROOT / "gemini-vault"
CHATGPT_DIR = ROOT / "chatgpt-vault"

CATEGORIES = [
    "社区",
    "数字制造",
    "交通工具",
    "UNNC",
    "团队管理",
    "工具",
    "产品分析",
    "FabLab",
    "autodesk",
    "HIL",
    "AI工具",
    "政策",
    "教育",
]

CHANNEL_LABELS = {
    "wechat": "WeChat",
    "web": "Web",
    "gemini": "Gemini",
    "chatgpt": "ChatGPT",
}

TEXT_EXTENSIONS = {
    ".md",
    ".txt",
    ".csv",
    ".json",
    ".html",
    ".htm",
    ".py",
    ".sh",
    ".ini",
    ".report",
    ".ipynb",
}

MEDIA_EXTENSIONS = {
    ".png",
    ".jpg",
    ".jpeg",
    ".gif",
    ".webp",
    ".svg",
    ".wav",
    ".mp3",
    ".mp4",
    ".mov",
    ".stl",
    ".kmz",
    ".key",
    ".ttf",
    ".ttc",
    ".zip",
    ".pyc",
    ".ldb",
    ".swp",
    ".bak",
}

SEARCH_TEXT_LIMIT = 5000
PAGE_TEXT_LIMIT = 60000
READ_BYTES_LIMIT = 8_000_000


class PlainTextHTMLParser(HTMLParser):
    """Small stdlib HTML-to-text parser tuned for Takeout exports."""

    BLOCK_TAGS = {
        "address",
        "article",
        "aside",
        "blockquote",
        "br",
        "dd",
        "div",
        "dl",
        "dt",
        "figcaption",
        "footer",
        "h1",
        "h2",
        "h3",
        "h4",
        "h5",
        "h6",
        "header",
        "hr",
        "li",
        "main",
        "nav",
        "ol",
        "p",
        "pre",
        "section",
        "table",
        "td",
        "th",
        "tr",
        "ul",
    }

    SKIP_TAGS = {"script", "style", "noscript", "svg"}

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self.skip_depth = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        tag = tag.lower()
        if tag in self.SKIP_TAGS:
            self.skip_depth += 1
            return
        if self.skip_depth:
            return
        if tag in self.BLOCK_TAGS:
            self.parts.append("\n")
        if tag == "a":
            href = dict(attrs).get("href")
            if href and not href.startswith("#"):
                self.parts.append(" ")
                self.parts.append(href)
                self.parts.append(" ")

    def handle_endtag(self, tag: str) -> None:
        tag = tag.lower()
        if tag in self.SKIP_TAGS and self.skip_depth:
            self.skip_depth -= 1
            return
        if self.skip_depth:
            return
        if tag in self.BLOCK_TAGS:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if self.skip_depth:
            return
        if data:
            self.parts.append(data)

    def text(self) -> str:
        return clean_text(" ".join(self.parts))


def now_iso() -> str:
    return datetime.now().astimezone().isoformat(timespec="seconds")


def clean_text(value: str | None) -> str:
    if not value:
        return ""
    value = value.replace("\u00a0", " ").replace("\ufeff", "")
    value = re.sub(r"[ \t\r\f\v]+", " ", value)
    value = re.sub(r"\n\s*", "\n", value)
    value = re.sub(r"\n{3,}", "\n\n", value)
    return value.strip()


def compact(value: str | None) -> str:
    return re.sub(r"\s+", " ", clean_text(value)).strip()


def truncate(value: str | None, limit: int) -> str:
    value = clean_text(value)
    if len(value) <= limit:
        return value
    return value[: limit - 1].rstrip() + "…"


def excerpt(value: str | None, limit: int = 260) -> str:
    return truncate(compact(value), limit)


def read_text_file(path: Path, max_bytes: int = READ_BYTES_LIMIT) -> str:
    try:
        data = path.read_bytes()
    except OSError:
        return ""
    if len(data) > max_bytes:
        head = data[: max_bytes // 2]
        tail = data[-max_bytes // 2 :]
        data = head + b"\n\n...[truncated]...\n\n" + tail
    for encoding in ("utf-8", "utf-8-sig", "gb18030", "latin-1"):
        try:
            return data.decode(encoding, errors="strict")
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="ignore")


def strip_html(value: str) -> str:
    parser = PlainTextHTMLParser()
    try:
        parser.feed(value)
        parser.close()
    except Exception:
        return compact(re.sub(r"<[^>]+>", " ", value))
    return parser.text()


def title_from_path(path: Path) -> str:
    name = path.name
    suffix = path.suffix
    if suffix:
        name = name[: -len(suffix)]
    name = re.sub(r"\s+", " ", name.replace("_", " ")).strip()
    return name or path.name


def path_href(path: str | Path) -> str:
    if isinstance(path, Path):
        path = path.as_posix()
    return quote(path, safe="/:#?&=%-_.~")


def stable_id(prefix: str, seed: str) -> str:
    return f"{prefix}-{hashlib.sha1(seed.encode('utf-8')).hexdigest()[:12]}"


def load_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def write_json(path: Path, data: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(data, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )


def write_js(path: Path, variable: str, data: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    payload = json.dumps(data, ensure_ascii=False, separators=(",", ":"))
    path.write_text(f"window.{variable} = {payload};\n", encoding="utf-8")


def document_search_text(*parts: str | None) -> str:
    return truncate("\n".join(p for p in parts if p), SEARCH_TEXT_LIMIT)


def normalize_category(category: str | None) -> str:
    category = compact(category)
    return category if category in CATEGORIES else "社区"


CATEGORY_KEYWORDS: dict[str, list[tuple[str, int]]] = {
    "autodesk": [
        ("autodesk", 10),
        ("fusion 360", 10),
        ("fusion360", 10),
        ("revit", 8),
        ("forge", 8),
        ("aps", 7),
        ("tinkercad", 7),
        ("alias", 6),
        ("pico gk", 5),
        ("picogk", 5),
    ],
    "HIL": [
        ("hil", 10),
        ("hardware in loop", 10),
        ("hardware-in-loop", 10),
        ("硬件在环", 10),
        ("云hil", 10),
        ("sim-to-real", 5),
        ("isaac sim", 4),
        ("isaac lab", 4),
    ],
    "UNNC": [
        ("unnc", 10),
        ("宁波诺丁汉", 10),
        ("nottingham", 8),
        ("delivery checklist", 5),
    ],
    "FabLab": [
        ("fablab", 10),
        ("fab lab", 10),
        ("fab academy", 10),
        ("fab foundation", 8),
        ("makerspace", 7),
        ("maker", 4),
        ("创客", 4),
        ("nexmaker", 8),
    ],
    "交通工具": [
        ("sophicar", 10),
        ("vehicle", 7),
        ("汽车", 7),
        ("小车", 7),
        ("移动机器人", 8),
        ("mobile robot", 8),
        ("robotics", 5),
        ("robot", 4),
        ("机器人", 5),
        ("autonomous", 6),
        ("无人驾驶", 6),
        ("自动驾驶", 6),
        ("底盘", 7),
        ("chassis", 7),
        ("bicycle", 5),
        ("pix moving", 8),
        ("drone", 5),
        ("无人机", 5),
        ("esp32", 4),
        ("电机", 4),
        ("a4988", 4),
    ],
    "数字制造": [
        ("digital fabrication", 10),
        ("数字制造", 10),
        ("additive", 7),
        ("3d printing", 8),
        ("3d打印", 8),
        ("manufacturing", 6),
        ("fabrication", 6),
        ("cad", 5),
        ("cam", 5),
        ("cnc", 6),
        ("dfm", 6),
        ("参数化", 6),
        ("parametric", 6),
        ("机械设计", 5),
        ("noyron", 6),
        ("leap71", 6),
        ("profabx", 8),
        ("加工", 4),
    ],
    "教育": [
        ("education", 8),
        ("课程", 8),
        ("教育", 8),
        ("教学", 7),
        ("夏令营", 9),
        ("camp", 5),
        ("student", 5),
        ("学生", 6),
        ("curriculum", 8),
        ("training", 5),
        ("pbl", 5),
        ("lesson", 4),
        ("高中", 4),
        ("中学生", 5),
        ("sophilab", 8),
        ("sophilabedu", 8),
    ],
    "AI工具": [
        ("ai工具", 10),
        ("ai 工具", 10),
        ("agent", 7),
        ("agents", 7),
        ("llm", 6),
        ("rag", 7),
        ("gemini", 6),
        ("claude", 7),
        ("codex", 6),
        ("deepseek", 6),
        ("qwen", 5),
        ("openai", 5),
        ("prompt", 5),
        ("提示词", 6),
        ("notebooklm", 8),
        ("generative ai", 6),
        ("mcp", 6),
        ("skill", 5),
        ("copilot", 4),
    ],
    "产品分析": [
        ("business plan", 10),
        ("商业计划", 10),
        ("bp", 5),
        ("pricing", 8),
        ("定价", 8),
        ("毛利", 8),
        ("market", 6),
        ("市场", 6),
        ("funding", 8),
        ("融资", 8),
        ("竞品", 8),
        ("competitor", 7),
        ("商业模式", 8),
        ("渠道", 5),
        ("pitch", 6),
    ],
    "团队管理": [
        ("团队", 8),
        ("管理", 5),
        ("workflow", 7),
        ("工作流", 7),
        ("process", 5),
        ("流程", 5),
        ("strategy", 5),
        ("战略", 5),
        ("chro", 8),
        ("招聘", 7),
        ("人才", 6),
        ("组织", 6),
        ("sop", 5),
    ],
    "政策": [
        ("policy", 8),
        ("政策", 10),
        ("政府", 8),
        ("高新区", 8),
        ("申报", 7),
        ("补贴", 7),
        ("项目申报", 8),
        ("教育部", 6),
        ("税", 4),
    ],
    "工具": [
        ("tool", 5),
        ("tools", 5),
        ("script", 5),
        ("代码", 4),
        ("python", 4),
        ("chrome", 5),
        ("maps", 4),
        ("drive", 4),
        ("glowgrow", 9),
        ("sensor", 4),
        ("传感器", 4),
        ("dashboard", 4),
        ("仪表盘", 4),
        ("app", 3),
    ],
    "社区": [
        ("community", 5),
        ("社区", 6),
        ("活动", 4),
        ("分享", 3),
    ],
}

TIE_PRIORITY = [
    "autodesk",
    "HIL",
    "UNNC",
    "FabLab",
    "交通工具",
    "数字制造",
    "教育",
    "AI工具",
    "产品分析",
    "团队管理",
    "政策",
    "工具",
    "社区",
]


def classify_gemini(title: str, path: Path, text: str) -> str:
    haystack = f"{title}\n{path.as_posix()}\n{text[:15000]}".lower()
    scores = Counter()

    for category, keywords in CATEGORY_KEYWORDS.items():
        for keyword, weight in keywords:
            count = haystack.count(keyword.lower())
            if count:
                scores[category] += min(count, 5) * weight

    path_text = path.as_posix().lower()
    if "sophicar" in path_text:
        scores["交通工具"] += 14
    if "sophilab" in path_text or "education" in path_text:
        scores["教育"] += 12
    if "google ai studio" in path_text:
        scores["AI工具"] += 8
    if "notebooklm" in path_text:
        scores["AI工具"] += 8
    if "nexpcb" in path_text:
        scores["数字制造"] += 8
    if "glowgrow" in path_text:
        scores["工具"] += 12
    if "chrome/" in path_text or "我的活动/maps" in path_text:
        scores["工具"] += 10

    if not scores:
        return "社区"
    best_score = max(scores.values())
    winners = [cat for cat, score in scores.items() if score == best_score]
    for category in TIE_PRIORITY:
        if category in winners:
            return category
    return winners[0]


def flatten_json(value: Any, limit: int = PAGE_TEXT_LIMIT) -> str:
    parts: list[str] = []

    def visit(node: Any, key: str = "", depth: int = 0) -> None:
        if len("\n".join(parts)) > limit:
            return
        if depth > 12:
            return
        if isinstance(node, dict):
            for child_key, child in node.items():
                if str(child_key).lower() in {"image", "thumbnail", "blob", "dataurl"}:
                    continue
                visit(child, str(child_key), depth + 1)
        elif isinstance(node, list):
            for child in node[:300]:
                visit(child, key, depth + 1)
        elif isinstance(node, (str, int, float, bool)) and node is not None:
            text = str(node)
            if len(text) > 4000:
                text = text[:4000] + "…"
            if text.startswith("data:image"):
                return
            if key and key.lower() in {"title", "name", "prompt", "text", "content", "description", "url"}:
                parts.append(f"{key}: {text}")
            elif len(text.strip()) > 1:
                parts.append(text)

    visit(value)
    return truncate(clean_text("\n".join(parts)), limit)


def extract_docx(path: Path) -> str:
    try:
        from docx import Document

        document = Document(str(path))
        parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables[:20]:
            for row in table.rows[:80]:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    parts.append(row_text)
        return truncate(clean_text("\n".join(parts)), PAGE_TEXT_LIMIT)
    except Exception as exc:
        return f"[docx extraction failed] {exc}"


def extract_pptx(path: Path) -> str:
    try:
        from pptx import Presentation

        prs = Presentation(str(path))
        parts: list[str] = []
        for index, slide in enumerate(prs.slides[:80], start=1):
            slide_parts = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    slide_parts.append(shape.text.strip())
            if slide_parts:
                parts.append(f"Slide {index}: " + "\n".join(slide_parts))
        return truncate(clean_text("\n\n".join(parts)), PAGE_TEXT_LIMIT)
    except Exception as exc:
        return f"[pptx extraction failed] {exc}"


def extract_xlsx(path: Path) -> str:
    try:
        from openpyxl import load_workbook

        workbook = load_workbook(str(path), read_only=True, data_only=True)
        parts: list[str] = []
        cell_budget = 3000
        for sheet in workbook.worksheets[:8]:
            parts.append(f"Sheet: {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                values = [str(cell) for cell in row if cell not in (None, "")]
                if values:
                    parts.append(" | ".join(values))
                cell_budget -= len(row)
                if cell_budget <= 0:
                    parts.append("[xlsx truncated]")
                    break
            if cell_budget <= 0:
                break
        workbook.close()
        return truncate(clean_text("\n".join(parts)), PAGE_TEXT_LIMIT)
    except Exception as exc:
        return f"[xlsx extraction failed] {exc}"


def extract_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                return "[pdf encrypted]"
        parts: list[str] = []
        for page_index, page in enumerate(reader.pages[:12], start=1):
            try:
                text = page.extract_text() or ""
            except Exception:
                text = ""
            if text.strip():
                parts.append(f"Page {page_index}\n{text}")
        return truncate(clean_text("\n\n".join(parts)), PAGE_TEXT_LIMIT)
    except Exception as exc:
        return f"[pdf extraction failed] {exc}"


def extract_zip_text(path: Path) -> str:
    try:
        with zipfile.ZipFile(path) as archive:
            parts: list[str] = []
            for info in archive.infolist()[:200]:
                inner = Path(info.filename)
                if info.is_dir():
                    continue
                if info.file_size > 1_000_000:
                    continue
                if inner.suffix.lower() not in TEXT_EXTENSIONS and not inner.name.endswith((".ts", ".tsx", ".js", ".jsx", ".css")):
                    continue
                try:
                    raw = archive.read(info)
                except Exception:
                    continue
                text = raw.decode("utf-8", errors="ignore")
                if inner.suffix.lower() in {".html", ".htm"}:
                    text = strip_html(text)
                parts.append(f"--- {info.filename} ---\n{text[:12000]}")
        return truncate(clean_text("\n\n".join(parts)), PAGE_TEXT_LIMIT)
    except Exception:
        return ""


def extract_text_by_type(path: Path) -> tuple[str, str]:
    suffix = path.suffix.lower()
    if suffix in {".html", ".htm"}:
        return strip_html(read_text_file(path)), "html"
    if suffix == ".json":
        try:
            return flatten_json(json.loads(read_text_file(path))), "json"
        except Exception:
            return read_text_file(path), "json"
    if suffix == ".ipynb":
        try:
            return flatten_json(json.loads(read_text_file(path))), "notebook"
        except Exception:
            return read_text_file(path), "notebook"
    if suffix in {".md", ".txt", ".csv", ".py", ".sh", ".ini", ".report"}:
        return read_text_file(path), suffix.lstrip(".")
    if suffix == ".docx":
        return extract_docx(path), "docx"
    if suffix == ".pptx":
        return extract_pptx(path), "pptx"
    if suffix == ".xlsx":
        return extract_xlsx(path), "xlsx"
    if suffix == ".pdf":
        return extract_pdf(path), "pdf"
    if not suffix:
        if zipfile.is_zipfile(path):
            return extract_zip_text(path), "ai-studio-package"
        raw = read_text_file(path)
        try:
            return flatten_json(json.loads(raw)), "json"
        except Exception:
            if raw.strip():
                return raw, "text"
    return "", suffix.lstrip(".") or "file"


def title_from_jsonish(path: Path, text: str) -> str:
    for pattern in (
        r'"title"\s*:\s*"([^"]{2,160})"',
        r'"name"\s*:\s*"([^"]{2,160})"',
        r"title:\s*([^\n]{2,160})",
        r"name:\s*([^\n]{2,160})",
    ):
        match = re.search(pattern, text)
        if match:
            return compact(match.group(1))
    return title_from_path(path)


def make_gemini_doc(
    *,
    title: str,
    source_path: Path,
    text: str,
    source_type: str,
    sequence: int = 0,
    date: str = "",
) -> dict[str, Any]:
    rel = source_path.relative_to(ROOT)
    doc_id = stable_id("gemini", f"{rel.as_posix()}::{title}::{sequence}")
    classification_path = source_path.relative_to(GEMINI_DIR) if source_path.is_relative_to(GEMINI_DIR) else rel
    category = classify_gemini(title, classification_path, text)
    body = truncate(text, PAGE_TEXT_LIMIT)
    return {
        "id": doc_id,
        "channel": "gemini",
        "channel_label": CHANNEL_LABELS["gemini"],
        "category": category,
        "title": truncate(title, 180),
        "href": f"gemini-vault/html/items/{doc_id}.html",
        "raw_href": path_href(rel),
        "source_path": rel.as_posix(),
        "source_url": "",
        "source_type": source_type,
        "date": date,
        "excerpt": excerpt(body),
        "search_text": document_search_text(title, rel.as_posix(), body),
        "text_chars": len(text),
        "body": body,
    }


def parse_gemini_activity(path: Path) -> list[dict[str, Any]]:
    html = read_text_file(path, max_bytes=40_000_000)
    chunks = re.split(r'<div class="outer-cell[^"]*"', html)
    docs: list[dict[str, Any]] = []
    for index, chunk in enumerate(chunks[1:], start=1):
        chunk = chunk.split('<div class="content-cell mdl-cell mdl-cell--12-col mdl-typography--caption">')[0]
        text = strip_html(chunk)
        text = re.sub(r"^\s*Gemini Apps\s*", "", text)
        text = clean_text(text)
        if len(text) < 30:
            continue
        date_match = re.search(r"(20\d{2}年\d{1,2}月\d{1,2}日\s+\d{1,2}:\d{2}:\d{2}\s+CST)", text)
        date = date_match.group(1) if date_match else ""
        lines = [compact(line) for line in text.splitlines() if compact(line)]
        prompt_line = ""
        for line in lines:
            if line.startswith("Prompted"):
                prompt_line = compact(line.replace("Prompted", "", 1))
                break
        if not prompt_line and lines:
            prompt_line = lines[0]
        title = prompt_line or f"Gemini Apps 活动 {index}"
        if date:
            title = f"{title} · {date}"
        docs.append(
            make_gemini_doc(
                title=title,
                source_path=path,
                text=text,
                source_type="gemini-activity",
                sequence=index,
                date=date,
            )
        )
    return docs


def parse_gemini_gems(path: Path) -> list[dict[str, Any]]:
    html = read_text_file(path, max_bytes=20_000_000)
    chunks = re.split(r"<b>名称：</b>", html)
    docs: list[dict[str, Any]] = []
    for index, chunk in enumerate(chunks[1:], start=1):
        fragment = "<b>名称：</b>" + chunk
        text = strip_html(fragment)
        text = clean_text(text)
        if len(text) < 30:
            continue
        name_match = re.search(r"名称：\s*([^\n]{1,120})", text)
        name = compact(name_match.group(1)) if name_match else f"Gemini Gem {index}"
        docs.append(
            make_gemini_doc(
                title=f"Gem: {name}",
                source_path=path,
                text=text,
                source_type="gemini-gem",
                sequence=index,
            )
        )
    return docs


def build_glowgrow_summary() -> dict[str, Any] | None:
    root = GEMINI_DIR / "Takeout 2" / "云端硬盘" / "glowgrow"
    if not root.exists():
        return None

    txt_files = [p for p in root.rglob("*.txt") if "data" in p.parts]
    report_files = list(root.rglob("*.report"))
    code_files = [p for p in root.iterdir() if p.suffix.lower() in {".py", ".sh", ".ini", ".md", ".txt"}]

    parts = [
        "glowgrow 项目资料摘要",
        f"传感器/运行日志 txt 文件：{len(txt_files)} 个",
        f"同步或处理 report 文件：{len(report_files)} 个",
        "这些日志以数据归档为主，已合并为一个摘要项，避免统一索引被逐条日志刷屏。",
    ]
    for file_path in code_files:
        text = read_text_file(file_path, max_bytes=120_000)
        if text.strip():
            parts.append(f"\n--- {file_path.name} ---\n{truncate(text, 8000)}")

    sample_names = [p.relative_to(root).as_posix() for p in (txt_files[:8] + report_files[:8])]
    if sample_names:
        parts.append("\n样例文件：\n" + "\n".join(sample_names))

    return make_gemini_doc(
        title="glowgrow 传感器数据与同步脚本归档",
        source_path=root / "readme2.md" if (root / "readme2.md").exists() else root,
        text=clean_text("\n".join(parts)),
        source_type="data-log-summary",
    )


def should_skip_gemini_path(path: Path) -> tuple[bool, str]:
    rel = path.relative_to(GEMINI_DIR)
    parts = rel.parts
    name = path.name.lower()
    suffix = path.suffix.lower()

    if name.startswith("."):
        return True, "hidden"
    if parts and parts[0] in {"html", "logs"}:
        return True, "generated"
    if name in {"search-index.json", "search-index.js"}:
        return True, "generated"
    if "__pycache__" in parts:
        return True, "cache"
    if "glowgrow" in parts and suffix in {".txt", ".report"} and ("data" in parts or "ossutil_output" in parts):
        return True, "aggregated-glowgrow-log"
    if suffix in MEDIA_EXTENSIONS:
        return True, suffix.lstrip(".") or "media"
    return False, ""


def build_gemini_docs() -> tuple[list[dict[str, Any]], dict[str, Any]]:
    docs: list[dict[str, Any]] = []
    stats: dict[str, Any] = {
        "skipped": Counter(),
        "errors": [],
        "source_type_counts": Counter(),
    }

    summary = build_glowgrow_summary()
    if summary:
        docs.append(summary)

    special_files = {
        GEMINI_DIR / "Takeout 3" / "我的活动" / "Gemini Apps" / "我的活动记录.html": parse_gemini_activity,
        GEMINI_DIR / "Takeout 3" / "Gemini" / "gemini_gems_data.html": parse_gemini_gems,
    }

    handled_specials: set[Path] = set()
    for path, parser in special_files.items():
        if path.exists():
            try:
                parsed = parser(path)
                docs.extend(parsed)
                handled_specials.add(path)
            except Exception as exc:
                stats["errors"].append({"path": path.relative_to(ROOT).as_posix(), "error": str(exc)})

    for path in sorted(GEMINI_DIR.rglob("*")):
        if not path.is_file():
            continue
        if path in handled_specials:
            continue
        skip, reason = should_skip_gemini_path(path)
        if skip:
            stats["skipped"][reason] += 1
            continue

        try:
            text, source_type = extract_text_by_type(path)
        except Exception as exc:
            stats["errors"].append({"path": path.relative_to(ROOT).as_posix(), "error": str(exc)})
            continue

        if not compact(text):
            stats["skipped"]["empty-text"] += 1
            continue

        title = title_from_jsonish(path, text) if source_type in {"json", "notebook", "ai-studio-package"} else title_from_path(path)
        docs.append(make_gemini_doc(title=title, source_path=path, text=text, source_type=source_type))

    deduped: list[dict[str, Any]] = []
    seen: set[str] = set()
    for doc in docs:
        key = f"{doc['source_path']}::{doc['title']}::{doc['source_type']}"
        if key in seen:
            continue
        seen.add(key)
        deduped.append(doc)
        stats["source_type_counts"][doc["source_type"]] += 1

    stats["skipped"] = dict(stats["skipped"])
    stats["source_type_counts"] = dict(stats["source_type_counts"])
    return deduped, stats


def load_wechat_docs() -> list[dict[str, Any]]:
    index_path = WECHAT_DIR / "search-index.json"
    if not index_path.exists():
        return []
    data = load_json(index_path)
    docs = []
    for item in data.get("documents", []):
        title = compact(item.get("title")) or "未命名微信资料"
        html_path = item.get("html_path") or ""
        md_path = item.get("md_path") or ""
        href = f"wechat-vault/{path_href(html_path or md_path)}" if (html_path or md_path) else "wechat-vault/html/index.html"
        source_url = item.get("final_url") or item.get("original_url") or ""
        doc_id = stable_id("wechat", f"{item.get('source_index')}::{title}::{html_path or md_path}")
        headings = "\n".join(item.get("headings") or [])
        text = item.get("text") or ""
        docs.append(
            {
                "id": doc_id,
                "channel": "wechat",
                "channel_label": CHANNEL_LABELS["wechat"],
                "category": normalize_category(item.get("category")),
                "title": truncate(title, 180),
                "href": href,
                "raw_href": source_url,
                "source_path": f"wechat-vault/{html_path or md_path}",
                "source_url": source_url,
                "source_type": "wechat-article",
                "date": item.get("published_at") or item.get("favorite_time") or "",
                "excerpt": excerpt(item.get("excerpt") or text),
                "search_text": document_search_text(title, item.get("account"), headings, item.get("excerpt"), text),
                "text_chars": item.get("text_chars") or len(text),
                "account": item.get("account") or "",
            }
        )
    return docs


def markdown_excerpt(path: Path) -> str:
    text = read_text_file(path, max_bytes=300_000)
    text = re.sub(r"!\[[^\]]*\]\([^)]+\)", " ", text)
    text = re.sub(r"\[[^\]]+\]\([^)]+\)", " ", text)
    text = re.sub(r"[#>*_`-]+", " ", text)
    return clean_text(text)


def load_web_docs() -> list[dict[str, Any]]:
    manifest = WEB_DIR / "web_pages_archive_manifest.csv"
    if not manifest.exists():
        return []
    docs: list[dict[str, Any]] = []
    with manifest.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        for row_index, row in enumerate(reader, start=1):
            title = compact(row.get("title")) or "未命名 Web 资料"
            html_path = row.get("html_path") or ""
            md_path = row.get("md_path") or ""
            href = f"web-vault/{path_href(html_path or md_path)}" if (html_path or md_path) else "web-vault/html/index.html"
            source_url = row.get("final_url") or row.get("url") or ""
            text = ""
            md_file = WEB_DIR / md_path if md_path else None
            if md_file and md_file.exists():
                text = markdown_excerpt(md_file)
            doc_id = stable_id("web", f"{row_index}::{title}::{source_url}")
            docs.append(
                {
                    "id": doc_id,
                    "channel": "web",
                    "channel_label": CHANNEL_LABELS["web"],
                    "category": normalize_category(row.get("category")),
                    "title": truncate(title, 180),
                    "href": href,
                    "raw_href": source_url,
                    "source_path": f"web-vault/{html_path or md_path}",
                    "source_url": source_url,
                    "source_type": "web-page",
                    "date": "",
                    "excerpt": excerpt(text) or compact(row.get("status")) or "Web 归档资料",
                    "search_text": document_search_text(title, source_url, row.get("status"), text),
                    "text_chars": int(row.get("text_chars") or 0),
                    "status": row.get("status") or "",
                    "http_status": row.get("http_status") or "",
                }
            )
    return docs


def strip_body(doc: dict[str, Any]) -> dict[str, Any]:
    clone = dict(doc)
    clone.pop("body", None)
    clone.pop("attachments", None)
    return clone


def summary_for_docs(docs: list[dict[str, Any]]) -> dict[str, Any]:
    channel_counts = Counter(doc["channel"] for doc in docs)
    category_counts = Counter(doc["category"] for doc in docs)
    matrix: dict[str, dict[str, int]] = {category: {} for category in CATEGORIES}
    for category in CATEGORIES:
        for channel in CHANNEL_LABELS:
            matrix[category][channel] = 0
    for doc in docs:
        matrix.setdefault(doc["category"], {}).setdefault(doc["channel"], 0)
        matrix[doc["category"]][doc["channel"]] += 1
    return {
        "generated_at": now_iso(),
        "document_count": len(docs),
        "channel_counts": dict(channel_counts),
        "category_counts": {category: category_counts.get(category, 0) for category in CATEGORIES},
        "category_channel_counts": matrix,
    }


def write_gemini_item_pages(docs: list[dict[str, Any]]) -> None:
    item_dir = GEMINI_DIR / "html" / "items"
    item_dir.mkdir(parents=True, exist_ok=True)
    for doc in docs:
        body = html_lib.escape(doc.get("body") or doc.get("excerpt") or "")
        title = html_lib.escape(doc["title"])
        category = html_lib.escape(doc["category"])
        source_type = html_lib.escape(doc.get("source_type") or "")
        source_path = html_lib.escape(doc.get("source_path") or "")
        date = html_lib.escape(doc.get("date") or "")
        raw = doc.get("raw_href") or ""
        raw_link = ""
        if raw.startswith("gemini-vault/"):
            raw_link = "../../" + raw[len("gemini-vault/") :]
        elif raw:
            raw_link = raw
        raw_anchor = f'<a href="{html_lib.escape(raw_link)}">打开原始文件</a>' if raw_link else ""
        html = f"""<!doctype html>
<html lang="zh-CN">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>{title} - Gemini Vault</title>
  <style>
    * {{ box-sizing: border-box; }}
    body {{ margin: 0; font: 15px/1.65 -apple-system, BlinkMacSystemFont, "Segoe UI", sans-serif; color: #1f2933; background: #f6f7f9; }}
    header {{ padding: 18px 22px; border-bottom: 1px solid #d9dee7; background: #fff; }}
    main {{ max-width: 980px; margin: 0 auto; padding: 22px; }}
    a {{ color: #0f766e; text-decoration: none; }}
    a:hover {{ text-decoration: underline; }}
    h1 {{ margin: 0 0 10px; font-size: 24px; line-height: 1.3; }}
    .meta {{ display: flex; flex-wrap: wrap; gap: 8px; color: #5f6b7a; font-size: 13px; }}
    .badge {{ border: 1px solid #cbd5df; border-radius: 8px; padding: 3px 8px; background: #fff; }}
    pre {{ white-space: pre-wrap; word-wrap: break-word; padding: 18px; background: #fff; border: 1px solid #d9dee7; border-radius: 8px; overflow: auto; }}
    nav {{ margin-top: 12px; display: flex; gap: 12px; flex-wrap: wrap; }}
  </style>
</head>
<body>
  <header>
    <h1>{title}</h1>
    <div class="meta">
      <span class="badge">Gemini</span>
      <span class="badge">{category}</span>
      <span class="badge">{source_type}</span>
      <span class="badge">{date}</span>
    </div>
    <nav>
      <a href="../../../index.html">返回 online-vault</a>
      <a href="../index.html">Gemini 索引</a>
      {raw_anchor}
    </nav>
  </header>
  <main>
    <p><strong>来源：</strong>{source_path}</p>
    <pre>{body}</pre>
  </main>
</body>
</html>
"""
        (item_dir / f"{doc['id']}.html").write_text(html, encoding="utf-8")


def root_index_html() -> str:
    return """<!doctype html>
<html lang="zh-CN">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>online-vault</title>
  <script src="online-vault-index.js"></script>
  <style>
    :root {
      --google-blue: #4285f4;
      --google-red: #db4437;
      --google-yellow: #f4b400;
      --google-green: #0f9d58;
      --ink: #202124;
      --muted: #5f6368;
      --line: #dadce0;
      --surface: #fff;
      --page: #f8fafd;
    }
    * { box-sizing: border-box; }
    body { margin: 0; font: 14px/1.5 -apple-system, BlinkMacSystemFont, "Segoe UI", sans-serif; color: var(--ink); background: var(--page); }
    a { color: var(--google-blue); text-decoration: none; }
    a:hover { text-decoration: underline; }
    header { position: sticky; top: 0; z-index: 5; background: var(--surface); border-bottom: 1px solid var(--line); }
    header::before { content: ""; display: block; height: 4px; background: linear-gradient(90deg, var(--google-blue) 0 25%, var(--google-red) 25% 50%, var(--google-yellow) 50% 75%, var(--google-green) 75% 100%); }
    .topbar { max-width: 1440px; margin: 0 auto; padding: 14px 18px; display: grid; grid-template-columns: 220px minmax(0, 1fr); align-items: center; gap: 14px; }
    h1 { margin: 0; font-size: 20px; line-height: 1.2; }
    .subtitle { color: var(--muted); font-size: 12px; margin-top: 3px; }
    .searchbar { display: grid; grid-template-columns: minmax(220px, 1fr) 160px 150px 140px; gap: 10px; align-items: center; }
    input, select { width: 100%; border: 1px solid var(--line); border-radius: 8px; background: #fff; color: var(--ink); padding: 9px 10px; font: inherit; }
    input:focus, select:focus { outline: 2px solid rgba(66, 133, 244, 0.18); border-color: var(--google-blue); }
    main { max-width: 1440px; margin: 0 auto; padding: 16px 18px 28px; }
    .dashboard { display: block; margin-bottom: 14px; }
    .dashboard-panel { background: #fff; border: 1px solid var(--line); border-radius: 8px; padding: 14px; box-shadow: 0 1px 2px rgba(60, 64, 67, 0.08); }
    .dashboard-wide { grid-column: 1 / -1; }
    .panel-head { display: flex; align-items: baseline; justify-content: space-between; gap: 12px; margin-bottom: 12px; }
    .panel-head h2 { margin: 0; font-size: 15px; line-height: 1.25; }
    .panel-note { color: var(--muted); font-size: 12px; }
    .dashboard-overview { margin-bottom: 16px; }
    .stats { display: grid; grid-template-columns: repeat(4, minmax(0, 1fr)); gap: 10px; }
    .stat { position: relative; display: block; overflow: hidden; background: #fff; border: 1px solid var(--line); border-radius: 8px; padding: 13px; min-height: 110px; color: inherit; }
    a.stat:hover { text-decoration: none; border-color: #c4d7fb; box-shadow: 0 2px 10px rgba(60, 64, 67, 0.12); transform: translateY(-1px); }
    a.stat:focus-visible { outline: 2px solid rgba(66, 133, 244, 0.35); outline-offset: 2px; }
    .stat::before { content: ""; position: absolute; inset: 0 auto 0 0; width: 4px; background: var(--accent, var(--google-blue)); }
    .stat .num { font-size: 24px; font-weight: 760; color: var(--ink); font-variant-numeric: tabular-nums; }
    .stat .label { color: var(--muted); font-size: 12px; margin-top: 2px; }
    .stat .share { color: var(--accent, var(--google-blue)); font-size: 12px; margin-top: 8px; font-weight: 650; }
    .stat .hint { color: #3c4043; font-size: 12px; margin-top: 10px; }
    .stat.wechat { --accent: var(--google-green); }
    .stat.web { --accent: var(--google-blue); }
    .stat.gemini { --accent: var(--google-yellow); }
    .stat.chatgpt { --accent: var(--google-red); }
    .matrix { display: grid; gap: 7px; }
    .matrix-row { display: grid; grid-template-columns: minmax(92px, 136px) minmax(0, 1fr) 66px; align-items: center; gap: 12px; width: 100%; border: 1px solid transparent; border-radius: 8px; background: transparent; padding: 7px 8px; color: inherit; font: inherit; text-align: left; cursor: pointer; }
    .matrix-row:hover { background: #f8fafd; border-color: #e8eaed; }
    .matrix-label { font-weight: 620; overflow: hidden; text-overflow: ellipsis; white-space: nowrap; }
    .stack { height: 12px; display: flex; overflow: hidden; border-radius: 999px; background: #edf0f4; }
    .seg { width: var(--width); min-width: 2px; }
    .seg.wechat { background: var(--google-green); }
    .seg.web { background: var(--google-blue); }
    .seg.gemini { background: var(--google-yellow); }
    .seg.chatgpt { background: var(--google-red); }
    .matrix-total { color: var(--muted); text-align: right; font-variant-numeric: tabular-nums; }
    .legend { display: flex; flex-wrap: wrap; gap: 10px; color: var(--muted); font-size: 12px; }
    .legend span { display: inline-flex; align-items: center; gap: 5px; }
    .dot { width: 8px; height: 8px; border-radius: 999px; background: var(--color); }
    .matrix-heading { display: flex; align-items: center; justify-content: space-between; gap: 12px; margin: 0 0 10px; }
    .matrix-heading h3 { margin: 0; font-size: 14px; }
    .layout { display: grid; grid-template-columns: 280px minmax(0, 1fr); gap: 14px; align-items: start; }
    aside { background: #fff; border: 1px solid var(--line); border-radius: 8px; padding: 12px; position: sticky; top: 80px; }
    aside h2 { margin: 0 0 10px; font-size: 14px; }
    .facet { display: grid; gap: 6px; margin-bottom: 14px; }
    .facet button { display: flex; justify-content: space-between; gap: 10px; width: 100%; border: 1px solid transparent; background: transparent; border-radius: 8px; padding: 7px 8px; color: #3c4043; font: inherit; cursor: pointer; text-align: left; }
    .facet button:hover, .facet button.active { background: #f1f5fe; border-color: #d2e3fc; }
    .count { color: #667085; font-variant-numeric: tabular-nums; }
    .result-head { display: flex; align-items: baseline; justify-content: space-between; gap: 12px; margin: 2px 0 10px; color: #667085; }
    .result-list { display: grid; gap: 9px; }
    .item { background: #fff; border: 1px solid var(--line); border-radius: 8px; padding: 13px 14px; }
    .item-title { font-size: 16px; font-weight: 680; line-height: 1.35; margin: 0 0 7px; }
    .meta { display: flex; flex-wrap: wrap; gap: 6px; margin-bottom: 7px; color: #667085; font-size: 12px; }
    .badge { border: 1px solid var(--line); border-radius: 8px; padding: 2px 7px; background: #fff; }
    .badge.wechat { border-color: #b7dfca; color: #137333; background: #e6f4ea; }
    .badge.web { border-color: #d2e3fc; color: #1967d2; background: #e8f0fe; }
    .badge.gemini { border-color: #fde293; color: #8a5f00; background: #fef7e0; }
    .badge.chatgpt { border-color: #f4b7ae; color: #a50e0e; background: #fce8e6; }
    .excerpt { color: #475467; margin: 0 0 8px; }
    .links { display: flex; flex-wrap: wrap; gap: 10px; font-size: 13px; }
    .empty { background: #fff; border: 1px dashed var(--line); border-radius: 8px; padding: 20px; color: #667085; }
    mark { background: #fff1a8; padding: 0 2px; border-radius: 3px; }
    @media (max-width: 980px) {
      .topbar { grid-template-columns: 1fr; }
      .searchbar { grid-template-columns: 1fr 1fr; }
      .layout { grid-template-columns: 1fr; }
      aside { position: static; }
      .stats { grid-template-columns: repeat(2, minmax(0, 1fr)); }
    }
    @media (max-width: 620px) {
      .searchbar, .stats { grid-template-columns: 1fr; }
      .matrix-row { grid-template-columns: 88px minmax(0, 1fr) 52px; gap: 8px; }
      main { padding: 12px; }
      .topbar { padding: 12px; }
    }
  </style>
</head>
<body>
  <header>
    <div class="topbar">
      <div>
        <h1>online-vault</h1>
        <div class="subtitle" id="generatedAt">正在载入索引</div>
      </div>
      <div class="searchbar" aria-label="筛选">
        <input id="queryInput" type="search" placeholder="搜索标题、摘要、路径或正文片段">
        <select id="categorySelect" aria-label="主题">
          <option value="">全部主题</option>
        </select>
        <select id="channelSelect" aria-label="渠道">
          <option value="">全部渠道</option>
          <option value="wechat">WeChat</option>
          <option value="web">Web</option>
          <option value="gemini">Gemini</option>
          <option value="chatgpt">ChatGPT</option>
        </select>
        <select id="sortSelect" aria-label="排序">
          <option value="relevance">相关优先</option>
          <option value="channel">渠道排序</option>
          <option value="category">主题排序</option>
          <option value="title">标题排序</option>
        </select>
      </div>
    </div>
  </header>

  <main>
    <section class="dashboard" aria-label="仪表盘">
      <div class="dashboard-panel dashboard-wide">
        <div class="panel-head">
          <h2>资料库仪表盘</h2>
          <span class="panel-note" id="dashboardUpdated">正在载入</span>
        </div>
        <div class="dashboard-overview">
          <div class="stats" aria-label="统计">
            <a class="stat wechat" href="wechat-vault/html/index.html"><div class="num" id="wechatCount">0</div><div class="label">WeChat</div><div class="share" id="wechatShare">0%</div><div class="hint">点击进入看板</div></a>
            <a class="stat web" href="web-vault/html/index.html"><div class="num" id="webCount">0</div><div class="label">Web</div><div class="share" id="webShare">0%</div><div class="hint">点击进入看板</div></a>
            <a class="stat gemini" href="gemini-vault/html/index.html"><div class="num" id="geminiCount">0</div><div class="label">Gemini</div><div class="share" id="geminiShare">0%</div><div class="hint">点击进入看板</div></a>
            <a class="stat chatgpt" href="chatgpt-vault/html/index.html"><div class="num" id="chatgptCount">0</div><div class="label">ChatGPT</div><div class="share" id="chatgptShare">0%</div><div class="hint">点击进入看板</div></a>
          </div>
        </div>
        <div class="matrix-heading">
          <h3>主题 × 渠道</h3>
          <div class="legend">
            <span><i class="dot" style="--color: var(--google-green)"></i>WeChat</span>
            <span><i class="dot" style="--color: var(--google-blue)"></i>Web</span>
            <span><i class="dot" style="--color: var(--google-yellow)"></i>Gemini</span>
            <span><i class="dot" style="--color: var(--google-red)"></i>ChatGPT</span>
          </div>
        </div>
        <div class="matrix" id="categoryMatrix"></div>
      </div>
    </section>

    <div class="layout">
      <aside>
        <h2>渠道</h2>
        <div class="facet" id="channelFacet"></div>
        <h2>主题</h2>
        <div class="facet" id="categoryFacet"></div>
      </aside>
      <section>
        <div class="result-head">
          <div id="resultSummary">0 条结果</div>
          <div id="limitNote"></div>
        </div>
        <div class="result-list" id="results"></div>
      </section>
    </div>
  </main>

  <script>
    const vault = window.ONLINE_VAULT_INDEX || { documents: [], categories: [], summary: {} };
    const docs = vault.documents || [];
    const categories = vault.categories || [];
    const channelLabels = { wechat: "WeChat", web: "Web", gemini: "Gemini", chatgpt: "ChatGPT" };
    const queryInput = document.getElementById("queryInput");
    const categorySelect = document.getElementById("categorySelect");
    const channelSelect = document.getElementById("channelSelect");
    const sortSelect = document.getElementById("sortSelect");
    const results = document.getElementById("results");
    const resultSummary = document.getElementById("resultSummary");
    const limitNote = document.getElementById("limitNote");
    const renderLimit = 350;

    function escapeHtml(value) {
      return String(value ?? "").replace(/[&<>"']/g, char => ({
        "&": "&amp;", "<": "&lt;", ">": "&gt;", '"': "&quot;", "'": "&#39;"
      }[char]));
    }

    function normalize(value) {
      return String(value || "").toLowerCase().replace(/\\s+/g, " ").trim();
    }

    function indexText(doc) {
      if (!doc._indexText) {
        doc._indexText = normalize(`${doc.title} ${doc.excerpt} ${doc.source_path} ${doc.search_text}`);
      }
      return doc._indexText;
    }

    function countBy(items, key) {
      return items.reduce((acc, item) => {
        acc[item[key]] = (acc[item[key]] || 0) + 1;
        return acc;
      }, {});
    }

    function formatNumber(value) {
      return Number(value || 0).toLocaleString();
    }

    function percent(value, total) {
      if (!total) return "0%";
      return `${Math.round((value / total) * 100)}%`;
    }

    function setStats() {
      const channelCounts = vault.summary?.channel_counts || countBy(docs, "channel");
      document.getElementById("wechatCount").textContent = formatNumber(channelCounts.wechat || 0);
      document.getElementById("webCount").textContent = formatNumber(channelCounts.web || 0);
      document.getElementById("geminiCount").textContent = formatNumber(channelCounts.gemini || 0);
      document.getElementById("chatgptCount").textContent = formatNumber(channelCounts.chatgpt || 0);
      document.getElementById("wechatShare").textContent = percent(channelCounts.wechat || 0, docs.length);
      document.getElementById("webShare").textContent = percent(channelCounts.web || 0, docs.length);
      document.getElementById("geminiShare").textContent = percent(channelCounts.gemini || 0, docs.length);
      document.getElementById("chatgptShare").textContent = percent(channelCounts.chatgpt || 0, docs.length);
      document.getElementById("dashboardUpdated").textContent = vault.generated_at ? `更新于 ${vault.generated_at}` : "索引已载入";
      document.getElementById("generatedAt").textContent = vault.generated_at ? `更新于 ${vault.generated_at}` : "索引已载入";
    }

    function setDashboard() {
      const matrix = vault.summary?.category_channel_counts || {};
      document.getElementById("categoryMatrix").innerHTML = categories.map(category => {
        const row = matrix[category] || {};
        const wechat = row.wechat || 0;
        const web = row.web || 0;
        const gemini = row.gemini || 0;
        const chatgpt = row.chatgpt || 0;
        const total = wechat + web + gemini + chatgpt;
        const safeTotal = Math.max(total, 1);
        return `<button type="button" class="matrix-row" data-category="${escapeHtml(category)}">
          <span class="matrix-label">${escapeHtml(category)}</span>
          <span class="stack">
            <span class="seg wechat" style="--width: ${(wechat / safeTotal) * 100}%"></span>
            <span class="seg web" style="--width: ${(web / safeTotal) * 100}%"></span>
            <span class="seg gemini" style="--width: ${(gemini / safeTotal) * 100}%"></span>
            <span class="seg chatgpt" style="--width: ${(chatgpt / safeTotal) * 100}%"></span>
          </span>
          <span class="matrix-total">${formatNumber(total)}</span>
        </button>`;
      }).join("");

      document.querySelectorAll("[data-category]").forEach(button => {
        button.addEventListener("click", () => {
          categorySelect.value = button.dataset.category || "";
          render();
        });
      });
    }

    function fillSelects() {
      categorySelect.innerHTML = '<option value="">全部主题</option>' + categories
        .map(category => `<option value="${escapeHtml(category)}">${escapeHtml(category)}</option>`)
        .join("");
    }

    function renderFacet(containerId, values, activeValue, onPick) {
      const container = document.getElementById(containerId);
      container.innerHTML = values.map(row => {
        const active = row.value === activeValue ? " active" : "";
        return `<button type="button" class="${active}" data-value="${escapeHtml(row.value)}"><span>${escapeHtml(row.label)}</span><span class="count">${row.count.toLocaleString()}</span></button>`;
      }).join("");
      container.querySelectorAll("button").forEach(button => {
        button.addEventListener("click", () => onPick(button.dataset.value || ""));
      });
    }

    function rank(doc, tokens) {
      if (!tokens.length) return 0;
      const title = doc._titleText || (doc._titleText = normalize(doc.title));
      const excerpt = doc._excerptText || (doc._excerptText = normalize(doc.excerpt));
      const haystack = indexText(doc);
      let score = 0;
      tokens.forEach(token => {
        if (title.includes(token)) score += 8;
        if (excerpt.includes(token)) score += 3;
        if (haystack.includes(token)) score += 1;
      });
      return score;
    }

    function filteredDocs(options = {}) {
      const category = options.ignoreCategory ? "" : categorySelect.value;
      const channel = options.ignoreChannel ? "" : channelSelect.value;
      const tokens = normalize(queryInput.value).split(" ").filter(Boolean);
      let rows = docs.filter(doc => {
        if (category && doc.category !== category) return false;
        if (channel && doc.channel !== channel) return false;
        if (!tokens.length) return true;
        return tokens.every(token => indexText(doc).includes(token));
      });

      const sort = sortSelect.value;
      if (sort === "title") {
        rows.sort((a, b) => a.title.localeCompare(b.title, "zh-Hans-CN"));
      } else if (sort === "category") {
        rows.sort((a, b) => a.category.localeCompare(b.category, "zh-Hans-CN") || a.title.localeCompare(b.title, "zh-Hans-CN"));
      } else if (sort === "channel") {
        rows.sort((a, b) => a.channel.localeCompare(b.channel) || a.category.localeCompare(b.category, "zh-Hans-CN"));
      } else {
        rows.sort((a, b) => rank(b, tokens) - rank(a, tokens));
      }
      return rows;
    }

    function render() {
      const rows = filteredDocs();
      const channelRows = filteredDocs({ ignoreChannel: true });
      const categoryRows = filteredDocs({ ignoreCategory: true });
      const channelCounts = countBy(channelRows, "channel");
      const categoryCounts = countBy(categoryRows, "category");
      renderFacet("channelFacet", [
        { value: "", label: "全部渠道", count: channelRows.length },
        { value: "wechat", label: "WeChat", count: channelCounts.wechat || 0 },
        { value: "web", label: "Web", count: channelCounts.web || 0 },
        { value: "gemini", label: "Gemini", count: channelCounts.gemini || 0 },
        { value: "chatgpt", label: "ChatGPT", count: channelCounts.chatgpt || 0 },
      ], channelSelect.value, value => { channelSelect.value = value; render(); });
      renderFacet("categoryFacet", [
        { value: "", label: "全部主题", count: categoryRows.length },
        ...categories.map(category => ({ value: category, label: category, count: categoryCounts[category] || 0 }))
      ], categorySelect.value, value => { categorySelect.value = value; render(); });

      resultSummary.textContent = `${rows.length.toLocaleString()} 条结果`;
      limitNote.textContent = rows.length > renderLimit ? `显示前 ${renderLimit.toLocaleString()} 条` : "";
      if (!rows.length) {
        results.innerHTML = '<div class="empty">没有匹配的资料。调整关键词、主题或渠道再试一次。</div>';
        return;
      }
      results.innerHTML = rows.slice(0, renderLimit).map(doc => {
        const raw = doc.raw_href && doc.raw_href !== doc.href ? `<a href="${escapeHtml(doc.raw_href)}">原链接/原文件</a>` : "";
        const chars = Number(doc.text_chars || 0).toLocaleString();
        const date = doc.date ? `<span class="badge">${escapeHtml(doc.date)}</span>` : "";
        return `<article class="item">
          <h2 class="item-title"><a href="${escapeHtml(doc.href)}">${escapeHtml(doc.title)}</a></h2>
          <div class="meta">
            <span class="badge ${escapeHtml(doc.channel)}">${escapeHtml(channelLabels[doc.channel] || doc.channel)}</span>
            <span class="badge">${escapeHtml(doc.category)}</span>
            <span class="badge">${escapeHtml(doc.source_type || "资料")}</span>
            ${date}
            <span class="badge">${chars} 字符</span>
          </div>
          <p class="excerpt">${escapeHtml(doc.excerpt || doc.source_path || "")}</p>
          <div class="links">
            <a href="${escapeHtml(doc.href)}">打开资料页</a>
            ${raw}
          </div>
        </article>`;
      }).join("");
    }

    [queryInput, categorySelect, channelSelect, sortSelect].forEach(control => control.addEventListener("input", render));
    setStats();
    setDashboard();
    fillSelects();
    render();
  </script>
</body>
</html>
"""


def gemini_index_html() -> str:
    return """<!doctype html>
<html lang="zh-CN">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>Gemini Vault</title>
  <script src="search-index.js"></script>
  <style>
    * { box-sizing: border-box; }
    body { margin: 0; font: 14px/1.55 -apple-system, BlinkMacSystemFont, "Segoe UI", sans-serif; color: #1f2933; background: #f6f7f9; }
    a { color: #0f766e; text-decoration: none; }
    a:hover { text-decoration: underline; }
    header { background: #fff; border-bottom: 1px solid #d8dee8; padding: 16px 18px; position: sticky; top: 0; z-index: 3; }
    .bar { max-width: 1280px; margin: 0 auto; display: grid; grid-template-columns: 220px 1fr auto; gap: 12px; align-items: center; }
    h1 { margin: 0; font-size: 20px; }
    .controls { display: grid; grid-template-columns: minmax(220px, 1fr) 160px 150px; gap: 10px; }
    input, select { border: 1px solid #cbd5df; border-radius: 8px; padding: 9px 10px; font: inherit; background: #fff; }
    nav { display: flex; gap: 8px; flex-wrap: wrap; justify-content: flex-end; }
    nav a { border: 1px solid #cbd5df; border-radius: 8px; padding: 7px 10px; color: #344054; }
    main { max-width: 1280px; margin: 0 auto; padding: 16px 18px 28px; }
    .summary { color: #667085; margin-bottom: 12px; }
    .list { display: grid; gap: 9px; }
    .item { background: #fff; border: 1px solid #d8dee8; border-radius: 8px; padding: 13px 14px; }
    .item h2 { margin: 0 0 7px; font-size: 16px; line-height: 1.35; }
    .meta { display: flex; gap: 6px; flex-wrap: wrap; color: #667085; font-size: 12px; margin-bottom: 7px; }
    .badge { border: 1px solid #cbd5df; border-radius: 8px; padding: 2px 7px; background: #fff; }
    .empty { background: #fff; border: 1px dashed #cbd5df; border-radius: 8px; padding: 20px; color: #667085; }
    @media (max-width: 880px) {
      .bar { grid-template-columns: 1fr; }
      .controls { grid-template-columns: 1fr; }
      nav { justify-content: flex-start; }
    }
  </style>
</head>
<body>
  <header>
    <div class="bar">
      <h1>Gemini Vault</h1>
      <div class="controls">
        <input id="query" type="search" placeholder="搜索 Gemini 资料">
        <select id="category"><option value="">全部主题</option></select>
        <select id="type"><option value="">全部类型</option></select>
      </div>
      <nav>
        <a href="../../index.html">online-vault</a>
        <a href="../search-index.json">JSON 索引</a>
      </nav>
    </div>
  </header>
  <main>
    <div class="summary" id="summary">正在载入</div>
    <div class="list" id="list"></div>
  </main>
  <script>
    const data = window.GEMINI_VAULT_INDEX || { documents: [], categories: [] };
    const docs = data.documents || [];
    const query = document.getElementById("query");
    const category = document.getElementById("category");
    const type = document.getElementById("type");
    const list = document.getElementById("list");
    const summary = document.getElementById("summary");
    const categories = data.categories || [];
    const types = [...new Set(docs.map(doc => doc.source_type).filter(Boolean))].sort();
    category.innerHTML = '<option value="">全部主题</option>' + categories.map(value => `<option>${escapeHtml(value)}</option>`).join("");
    type.innerHTML = '<option value="">全部类型</option>' + types.map(value => `<option>${escapeHtml(value)}</option>`).join("");

    function escapeHtml(value) {
      return String(value ?? "").replace(/[&<>"']/g, char => ({
        "&": "&amp;", "<": "&lt;", ">": "&gt;", '"': "&quot;", "'": "&#39;"
      }[char]));
    }
    function normalize(value) { return String(value || "").toLowerCase().replace(/\\s+/g, " ").trim(); }
    function render() {
      const tokens = normalize(query.value).split(" ").filter(Boolean);
      let rows = docs.filter(doc => {
        if (category.value && doc.category !== category.value) return false;
        if (type.value && doc.source_type !== type.value) return false;
        if (!tokens.length) return true;
        const haystack = normalize(`${doc.title} ${doc.excerpt} ${doc.source_path} ${doc.search_text}`);
        return tokens.every(token => haystack.includes(token));
      });
      summary.textContent = `${rows.length.toLocaleString()} / ${docs.length.toLocaleString()} 条 Gemini 资料，更新于 ${data.generated_at || ""}`;
      if (!rows.length) {
        list.innerHTML = '<div class="empty">没有匹配的 Gemini 资料。</div>';
        return;
      }
      list.innerHTML = rows.slice(0, 350).map(doc => `
        <article class="item">
          <h2><a href="${escapeHtml(doc.href.replace(/^gemini-vault\\/html\\//, ""))}">${escapeHtml(doc.title)}</a></h2>
          <div class="meta">
            <span class="badge">${escapeHtml(doc.category)}</span>
            <span class="badge">${escapeHtml(doc.source_type)}</span>
            <span class="badge">${Number(doc.text_chars || 0).toLocaleString()} 字符</span>
          </div>
          <p>${escapeHtml(doc.excerpt || "")}</p>
        </article>
      `).join("");
    }
    [query, category, type].forEach(control => control.addEventListener("input", render));
    render();
  </script>
</body>
</html>
"""


def main() -> None:
    generated_at = now_iso()
    wechat_docs = load_wechat_docs()
    web_docs = load_web_docs()
    gemini_docs, gemini_stats = build_gemini_docs()

    try:
        from build_chatgpt_vault import build_chatgpt_docs, write_chatgpt_outputs

        chatgpt_docs, chatgpt_stats = build_chatgpt_docs(ROOT)
        write_chatgpt_outputs(ROOT, generated_at, chatgpt_docs, chatgpt_stats)
    except Exception as exc:
        print(f"[warn] chatgpt export rendering failed: {exc}")
        chatgpt_docs = []
        chatgpt_stats = {"errors": [str(exc)]}

    write_gemini_item_pages(gemini_docs)

    gemini_public_docs = [strip_body(doc) for doc in gemini_docs]
    gemini_index = {
        "generated_at": generated_at,
        "categories": CATEGORIES,
        "summary": summary_for_docs(gemini_public_docs),
        "documents": gemini_public_docs,
    }
    write_json(GEMINI_DIR / "search-index.json", gemini_index)
    write_js(GEMINI_DIR / "html" / "search-index.js", "GEMINI_VAULT_INDEX", gemini_index)

    chatgpt_public_docs = [strip_body(doc) for doc in chatgpt_docs]

    all_docs = wechat_docs + web_docs + gemini_public_docs + chatgpt_public_docs
    online_index = {
        "generated_at": generated_at,
        "categories": CATEGORIES,
        "channels": CHANNEL_LABELS,
        "summary": summary_for_docs(all_docs),
        "documents": all_docs,
    }
    write_json(ROOT / "online-vault-index.json", online_index)
    write_js(ROOT / "online-vault-index.js", "ONLINE_VAULT_INDEX", online_index)
    (ROOT / "index.html").write_text(root_index_html(), encoding="utf-8")

    try:
        from render_channel_dashboards import render_chatgpt_dashboard, render_gemini_dashboard, render_web_dashboard

        render_web_dashboard(ROOT)
        render_gemini_dashboard(ROOT)
        render_chatgpt_dashboard(ROOT)
    except Exception as exc:
        print(f"[warn] channel dashboard rendering failed: {exc}")
        (GEMINI_DIR / "html" / "index.html").write_text(gemini_index_html(), encoding="utf-8")

    online_summary = summary_for_docs(all_docs)
    online_summary["gemini_source_type_counts"] = gemini_stats["source_type_counts"]
    online_summary["gemini_skipped"] = gemini_stats["skipped"]
    online_summary["gemini_errors"] = gemini_stats["errors"][:80]
    online_summary["chatgpt_model_counts"] = chatgpt_stats.get("model_counts", {})
    online_summary["chatgpt_role_counts"] = chatgpt_stats.get("role_counts", {})
    online_summary["chatgpt_errors"] = chatgpt_stats.get("errors", [])[:80]
    write_json(ROOT / "logs" / "online-vault-summary.json", online_summary)
    write_json(
        GEMINI_DIR / "logs" / "reindex-summary.json",
        {
            "generated_at": generated_at,
            "document_count": len(gemini_public_docs),
            "category_counts": summary_for_docs(gemini_public_docs)["category_counts"],
            "source_type_counts": gemini_stats["source_type_counts"],
            "skipped": gemini_stats["skipped"],
            "errors": gemini_stats["errors"][:80],
        },
    )

    print(json.dumps(online_summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
