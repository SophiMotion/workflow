"""Export a separate Want workbook; never writes the source personnel workbook."""
import json,sys,re
from datetime import datetime
from zoneinfo import ZoneInfo
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter
src,dst=sys.argv[1:3]
data=json.load(open(src)); rows=data['rows']; wb=Workbook();wb.remove(wb.active)
HEADER='254B80'; GREEN='E6F4EA'
def value(cell,v):
    if isinstance(v,str):
        cell.value=re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]','',v)[:32767];cell.data_type='s'
    else: cell.value=v

def stamp(v):
    if not v:return None
    try:return datetime.fromisoformat(v.replace('Z','+00:00')).astimezone(ZoneInfo('Asia/Shanghai')).replace(tzinfo=None)
    except (ValueError,TypeError):return v

def sheet(name,headers,records,widths):
    ws=wb.create_sheet(name);ws.append(headers);ws.freeze_panes='A2';ws.sheet_view.showGridLines=False
    for j,title in enumerate(headers,1):
        c=ws.cell(1,j);c.font=Font(name='Arial',size=11,bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor=HEADER);c.alignment=Alignment(vertical='center',wrap_text=True)
        ws.column_dimensions[get_column_letter(j)].width=widths[j-1]
    ws.row_dimensions[1].height=30
    for i,record in enumerate(records,2):
        for j,v in enumerate(record,1):
            c=ws.cell(i,j);value(c,v);c.font=Font(name='Arial',size=11,color='26364B');c.alignment=Alignment(vertical='top',wrap_text=True)
            if isinstance(v,datetime):c.number_format='yyyy-mm-dd hh:mm'
        ws.row_dimensions[i].height=75
    ws.auto_filter.ref=ws.dimensions
    return ws
person=sheet('人员分类',['记录 ID','岗位','姓名','学校','专业','技能','项目','阅读层级','匹配分类','分类来源','职业依据','待核实项','沟通状态','首次阅读','最后复核','资料来源'],[[r.get(k,'') if k not in ('reviewedAt','lastReviewAt') else stamp(r.get(k)) for k in ['id','roleName','name','school','major','skills','projects','reviewLevel','category','categorySource','evidence','gaps','communicationStatus','reviewedAt','lastReviewAt','sourceUrl']] for r in rows],[35,28,16,28,22,50,55,16,25,20,55,45,22,22,22,40])
review=sheet('复核记录',['记录 ID','姓名','岗位','时间','分类','复核意见'],[[r['id'],r.get('name'),r.get('roleName'),stamp(h['at']),h['category'],h['note']] for r in rows for h in r.get('reviewHistory',[])],[35,16,28,22,25,85])
messages=[r for r in rows if r.get('message') or r.get('reply')]
msg=sheet('消息记录',['记录 ID','姓名','岗位','状态','发送时间','消息正文','核验依据','反馈摘要','反馈范围'],[[r['id'],r.get('name'),r.get('roleName'),r['communicationStatus'],stamp(r.get('sentAt')),r.get('message'),r.get('verification'),r.get('reply'),r.get('replySource')] for r in messages],[35,16,28,22,22,85,50,60,45])
for i,r in enumerate(messages,2):
    if r.get('status')=='已发送' and all(r.get(k) for k in ['sentAt','message','verification']):
        for j in [4,5,6]:msg.cell(i,j).fill=PatternFill('solid',fgColor=GREEN)
model=sheet('模型评估',['记录 ID','姓名','岗位','时间','分类','匹配依据','待核实项','资料已变化'],[[r['id'],r.get('name'),r.get('roleName'),stamp(r['assessment'].get('at')),r['assessment'].get('category'),r['assessment'].get('evidence'),r['assessment'].get('gaps'),'是' if r.get('assessmentStale') else '否'] for r in rows if r.get('assessment')],[35,16,28,22,25,70,60,18])
sheet('导出说明',['项目','内容'],[['导出记录数',len(rows)],['生成时间',stamp(data['meta']['generatedAt'])],['筛选条件',json.dumps(data['meta'].get('filters',{}),ensure_ascii=False)],['口径','本地历史资料；摘要与完整简历核查分开，复核不会改写首次阅读或发送时间。'],['高亮','消息记录中绿色底色表示实际已发送且存在核验依据；不代表录用或面试结果。']],[25,100])
wb.save(dst)
